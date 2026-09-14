#!/usr/bin/env python3
"""Gera minuta HTML e, quando python-docx existir, DOCX do contrato."""

from __future__ import annotations

import argparse
import html
import json
import sys
from pathlib import Path


REQUIRED = ("prestador_nome", "prestador_documento", "prestador_endereco", "cliente_nome", "cliente_documento", "cliente_endereco", "servico", "valor", "pagamento", "prazo", "foro")


def text(v) -> str:
    return html.escape(str(v or "(preencher)"))


def clauses(d: dict) -> list[tuple[str, str]]:
    maintenance = "Não contratada." if not d.get("manutencao") else f"Contratada por {d.get('valor_manutencao', '(preencher)')} mensais, limitada ao escopo descrito."
    return [
        ("Objeto", str(d["servico"])),
        ("Escopo", str(d.get("escopo") or "Criação ou redesign de site institucional e publicação, conforme proposta aprovada. Pedidos adicionais serão orçados separadamente.")),
        ("Prazo", str(d["prazo"])),
        ("Preço e pagamento", f"Valor total de {d['valor']}, pago da seguinte forma: {d['pagamento']}."),
        ("Responsabilidades do cliente", "Fornecer textos, imagens, logotipos, acessos e aprovações necessários, garantindo que possui autorização para uso desses materiais."),
        ("Domínio e hospedagem", str(d.get("dominio_hospedagem") or "Domínio e serviços contratados em nome do cliente, salvo previsão expressa em proposta.")),
        ("Manutenção", maintenance),
        ("Limites técnicos", "O objeto é um site institucional. Sistemas com login, banco de dados, pagamento, prontuário ou dados sensíveis não estão incluídos."),
        ("Rescisão", str(d.get("rescisao") or "A rescisão e os valores devidos observarão as entregas já realizadas e o acordo escrito entre as partes.")),
        ("Foro", f"Fica eleito o foro de {d['foro']}, sem prejuízo de regra legal obrigatória."),
    ]


def build_html(d: dict) -> str:
    body = "".join(f"<h2>{i+1}. {text(title)}</h2><p>{text(content)}</p>" for i, (title, content) in enumerate(clauses(d)))
    return f'''<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Contrato — {text(d['cliente_nome'])}</title><style>@page{{size:A4;margin:2cm}}body{{font:12pt/1.55 Georgia,serif;color:#111;max-width:18cm;margin:2cm auto}}h1{{font-size:18pt;text-align:center}}h2{{font-size:12pt;margin-top:20px}}p{{text-align:justify}}.sign{{display:grid;grid-template-columns:1fr 1fr;gap:40px;margin-top:70px}}.line{{border-top:1px solid;padding-top:8px;text-align:center}}small{{display:block;margin-top:50px;color:#666}}@media print{{body{{margin:0}}}}</style></head><body><h1>CONTRATO DE PRESTAÇÃO DE SERVIÇOS</h1><p><b>PRESTADORA:</b> {text(d['prestador_nome'])}, {text(d['prestador_documento'])}, {text(d['prestador_endereco'])}.</p><p><b>CLIENTE:</b> {text(d['cliente_nome'])}, {text(d['cliente_documento'])}, {text(d['cliente_endereco'])}.</p>{body}<div class="sign"><div class="line">{text(d['prestador_nome'])}<br>Prestadora</div><div class="line">{text(d['cliente_nome'])}<br>Cliente</div></div><p>Data: {text(d.get('data'))}</p><small>Minuta-base operacional. Recomenda-se revisão por profissional jurídico antes da assinatura.</small></body></html>'''


def build_docx(d: dict, target: Path) -> bool:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return False
    doc = Document(); title = doc.add_heading("CONTRATO DE PRESTAÇÃO DE SERVIÇOS", 0); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"PRESTADORA: {d['prestador_nome']}, {d['prestador_documento']}, {d['prestador_endereco']}.")
    doc.add_paragraph(f"CLIENTE: {d['cliente_nome']}, {d['cliente_documento']}, {d['cliente_endereco']}.")
    for i, (name, content) in enumerate(clauses(d), 1):
        doc.add_heading(f"{i}. {name}", level=2); doc.add_paragraph(content)
    doc.add_paragraph("\n\n________________________________\n" + str(d["prestador_nome"]) + " — Prestadora")
    doc.add_paragraph("\n________________________________\n" + str(d["cliente_nome"]) + " — Cliente")
    doc.add_paragraph("Data: " + str(d.get("data") or "(preencher)"))
    doc.add_paragraph("Minuta-base operacional. Recomenda-se revisão por profissional jurídico antes da assinatura.")
    doc.save(target); return True


def main() -> int:
    p = argparse.ArgumentParser(); p.add_argument("--dados", required=True); p.add_argument("--workspace", required=True); p.add_argument("--slug", required=True); args = p.parse_args()
    d = json.loads(Path(args.dados).read_text(encoding="utf-8")); missing = [k for k in REQUIRED if not d.get(k)]
    if missing:
        raise ValueError("Campos obrigatórios ausentes: " + ", ".join(missing))
    folder = Path(args.workspace).expanduser().resolve() / "sites" / args.slug; folder.mkdir(parents=True, exist_ok=True)
    (folder / "dados-contrato.json").write_text(json.dumps(d, ensure_ascii=False, indent=2), encoding="utf-8")
    (folder / "contrato.html").write_text(build_html(d), encoding="utf-8")
    made = build_docx(d, folder / "contrato.docx")
    print(json.dumps({"html": str(folder / 'contrato.html'), "docx": str(folder / 'contrato.docx') if made else None, "aviso": None if made else "Instale python-docx para gerar DOCX"}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERRO: {exc}", file=sys.stderr); raise SystemExit(1)
